from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/nhatanhly/Belucar_app/docs/Dac_ta_flow_dat_ve_xe_concert.docx")

# Preset: compact_reference_guide, header pattern: memo_masthead.
PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

GREEN = "0B4F3D"
GREEN_DARK = "073B2E"
GOLD = "D5A62A"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
INK = "202624"
MUTED = "5E6864"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "EAF4F0"
LIGHT_GOLD = "FFF7DA"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "9B1C1C"
ACTIVE_NUM_ID = None


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, color=INK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in CELL_MARGIN_DXA.items():
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_WIDTH_DXA, (widths, sum(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="C7CFCC", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def set_paragraph_spacing(paragraph, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_text(doc, text, bold=False, color=INK, italic=False, after=6, size=11):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_label_text(doc, label, text, after=5):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, after=after)
    r = p.add_run(label + ": ")
    set_run_font(r, size=11, color=INK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.375 if level == 0 else 0.625)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(" " + text)
    set_run_font(r, size=11, color=INK)
    return p


def create_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    style_num_id = int(
        doc.styles["List Number"]._element.pPr.numPr.numId.val
    )
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == style_num_id
    )
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    new_num_id = max(existing) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_num_id


def start_numbered_sequence(doc):
    global ACTIVE_NUM_ID
    ACTIVE_NUM_ID = create_numbering_instance(doc)


def set_paragraph_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num = num_pr.find(qn("w:numId"))
    if num is None:
        num = OxmlElement("w:numId")
        num_pr.append(num)
    num.set(qn("w:val"), str(num_id))


def add_numbered(doc, title, description):
    p = doc.add_paragraph(style="List Number")
    if ACTIVE_NUM_ID is not None:
        set_paragraph_numbering(p, ACTIVE_NUM_ID)
    p.paragraph_format.left_indent = Inches(0.375)
    p.paragraph_format.first_line_indent = Inches(-0.188)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(title + ". ")
    set_run_font(r, size=11, color=GREEN_DARK, bold=True)
    r = p.add_run(description)
    set_run_font(r, size=11, color=INK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_callout(doc, label, text, fill=LIGHT_GOLD, accent=GOLD):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, after=0)
    r = p.add_run(label + ": ")
    set_run_font(r, size=11, color=GREEN_DARK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for idx, value in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0, line=1.1)
        r = p.add_run(value)
        set_run_font(r, size=9.2, color=GREEN_DARK, bold=True)

    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, after=0, line=1.12)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK, bold=(idx == 0))
    set_table_geometry(table, widths)
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(2)
    return table


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# Style token map from compact_reference_guide.
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

heading_tokens = {
    1: (16, BLUE, 18, 10),
    2: (13, BLUE, 14, 7),
    3: (12, BLUE_DARK, 10, 5),
}
for level, (size, color, before, after) in heading_tokens.items():
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = rgb(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for list_style_name in ("List Bullet", "List Bullet 2", "List Number"):
    style = doc.styles[list_style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.25

# Running header and footer.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hp.paragraph_format.space_after = Pt(0)
hr = hp.add_run("BELUCAR | ĐẶC TẢ FLOW CONCERT")
set_run_font(hr, size=8.5, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_after = Pt(0)
fr = fp.add_run("Nội bộ - Trang ")
set_run_font(fr, size=9, color=MUTED)
add_page_field(fp)

# Memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(10)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("ĐẶC TẢ FLOW ĐẶT VÉ XE CONCERT")
set_run_font(r, size=23, color=GREEN_DARK, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(15)
r = p.add_run("Tổng hợp nghiệp vụ, logic ngày đi - ngày về và hướng tích hợp API")
set_run_font(r, size=13.5, color=MUTED, bold=True)

add_label_text(doc, "Dự án", "Belucar - Vé xe đi concert")
add_label_text(doc, "Ngày tổng hợp", "14/08/2026")
add_label_text(doc, "Đối tượng đọc", "Mobile, Backend, QA, Product")
add_label_text(doc, "Trạng thái", "Bản demo UI/logic, chưa tích hợp API tạo và thanh toán vé")

rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(7)
rule.paragraph_format.space_after = Pt(12)
p_pr = rule._p.get_or_add_pPr()
p_bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:space"), "1")
bottom.set(qn("w:color"), GOLD)
p_bdr.append(bottom)
p_pr.append(p_bdr)

add_callout(
    doc,
    "Nguyên tắc cốt lõi",
    "Mọi lựa chọn trên UI phải được chuẩn hóa thành các lượt xe cụ thể. Backend là nguồn dữ liệu chính thức cho giá, trạng thái vé và mã QR; mobile chỉ hiển thị kết quả và kiểm tra hợp lệ cơ bản.",
)

add_heading(doc, "1. Mục tiêu và phạm vi", 1)
add_text(
    doc,
    "Tài liệu mô tả flow đã thống nhất trong phiên làm việc, các case ngày đi/ngày về, cách tính giá hiện tại, hành vi của vé và đề xuất hợp đồng API. Phạm vi hiện tại là giao diện mô phỏng để duyệt nghiệp vụ; chưa có API tạo đơn, thanh toán hoặc kiểm vé phía tài xế.",
)
add_bullet(doc, "Sự kiện có hai ngày phục vụ: 24/10/2026 và 25/10/2026.")
add_bullet(doc, "Điểm đến của lượt đi được cố định tại Sân vận động Quốc gia Mỹ Đình.")
add_bullet(doc, "Giá cơ sở trong demo: 120.000 đồng/khách/lượt.")
add_bullet(doc, "Ba loại hành trình: Vé lượt đi, Vé lượt về và Khứ hồi.")
add_bullet(doc, "Người dùng có thể có nhiều lần mua; mỗi lần mua tạo một vé/đơn riêng, không ghi đè vé cũ.")

add_heading(doc, "2. Tổng quan flow người dùng", 1)
add_heading(doc, "2.1. Khách hàng đã đăng nhập", 2)
start_numbered_sequence(doc)
add_numbered(doc, "Điểm vào", "Mở từ banner Home, ô dịch vụ concert ở đầu Home hoặc mục điều hướng concert trên bottom bar.")
add_numbered(doc, "Chọn thao tác", "Mua vé mới hoặc xem vé đã có.")
add_numbered(doc, "Tạo vé", "Chọn hành trình, điểm đón/điểm trả, ngày, giờ (nếu có), loại vé, số lượng và xác nhận.")
add_numbered(doc, "Nhận vé", "Sau khi tạo thành công, mở ngay màn hình chi tiết vé có QR để tài xế kiểm vé.")
add_numbered(doc, "Xem lại", "Nếu tài khoản có một vé thì mở thẳng chi tiết; nếu có từ hai vé trở lên thì mở danh sách để chọn.")

add_heading(doc, "2.2. Khách hàng chưa đăng nhập", 2)
start_numbered_sequence(doc)
add_numbered(doc, "Điểm vào", "Tại màn hình đăng nhập, giữ nút Mua vé đi concert bên cạnh nút Đăng nhập.")
add_numbered(doc, "Chế độ", "Trong màn hình concert có công tắc Mua vé/Tra vé. Tra vé dùng số điện thoại đã mua.")
add_numbered(doc, "Thông tin bắt buộc", "Email và số điện thoại; số điện thoại đồng thời là tài khoản đăng nhập.")
add_numbered(doc, "Tạo tài khoản", "Bản demo sinh username và mật khẩu tạm. Khi có API, server phải cấp mật khẩu/thông tin kích hoạt.")
add_numbered(doc, "Bàn giao vé", "Hiển thị vé vừa mua, QR và ghi chú yêu cầu đăng nhập bằng thông tin được cấp để xuất trình vé khi lên xe.")
add_numbered(doc, "Ghi nhớ đăng nhập", "Số điện thoại và mật khẩu phiên cuối được lưu an toàn để tự điền lại khi token hết hạn hoặc sau khi đăng xuất; người dùng có thể bấm X để xóa cả hai.")

add_heading(doc, "3. Dữ liệu đầu vào và ràng buộc UI", 1)
add_table(
    doc,
    ["Nhóm", "Quy tắc hiện tại"],
    [
        ("Ngày concert", "Chỉ cho chọn 24/10/2026, 25/10/2026 hoặc cả hai ngày."),
        ("Loại hành trình", "Ba lựa chọn ngang: Vé lượt đi, Vé lượt về, Khứ hồi."),
        ("Vị trí", "Lượt đi: người dùng chọn điểm đón, điểm đến cố định Mỹ Đình. Lượt về: điểm đón là Mỹ Đình, người dùng chọn điểm trả."),
        ("Khung giờ", "Lượt đi/khứ hồi chọn 14:00, 15:30, 17:00 hoặc 18:30. Lượt về không chọn giờ."),
        ("Loại vé", "Ghế lẻ, xe 5 chỗ hoặc xe 7 chỗ. Xe 5/7 chỗ có tùy chọn Bao xe."),
        ("Số lượng", "Tối thiểu 1. Ghế lẻ tối đa 6; xe 5 chỗ tối đa 5; xe 7 chỗ tối đa 7. Bao xe khóa số lượng theo sức chứa."),
        ("Khách mới", "Email hợp lệ; số điện thoại gồm 10 số và bắt đầu bằng 0."),
    ],
    [1900, 7460],
    font_size=9.6,
)

add_heading(doc, "4. Mô hình trạng thái nghiệp vụ", 1)
add_text(doc, "Mobile hiện dùng các trạng thái sau để điều khiển giao diện và giá:")
add_table(
    doc,
    ["Trạng thái", "Ý nghĩa"],
    [
        ("selectedDates", "Danh sách ngày đi đã chọn. Luôn còn ít nhất một ngày trong chế độ lượt đi/khứ hồi."),
        ("selectedReturnDates", "Danh sách ngày về; trong chế độ chỉ về đây là ngày phục vụ chính."),
        ("isReturnOnly", "Chỉ mua lượt về; không có ngày đi và không yêu cầu khung giờ."),
        ("wantsReturnTrip", "Người dùng bật ô Đặt lượt về trong flow lượt đi."),
        ("isRoundTrip", "Ngày đi và ngày về trùng nhau hoàn toàn; áp dụng giảm 10% cho từng cặp lượt."),
        ("isOvernightJourney", "Chọn đi cả 24 và 25 nhưng chỉ về 25; hệ thống chuẩn hóa còn đi 24 và về 25."),
        ("isCharter", "Bao toàn bộ xe 5/7 chỗ; số ghế tính tiền bằng sức chứa xe."),
    ],
    [2200, 7160],
    font_size=9.5,
)

add_callout(
    doc,
    "Khuyến nghị kỹ thuật",
    "Khi nối API, không nên tiếp tục suy luận hành trình chỉ từ nhiều cờ Boolean. Hãy chuẩn hóa thành journeyType và danh sách legs để mỗi lượt đi/lượt về có ngày, điểm đón, điểm trả và chính sách thời gian riêng.",
    fill=LIGHT_GREEN,
    accent=GREEN,
)

add_heading(doc, "5. Ma trận case ngày đi - ngày về", 1)
add_text(
    doc,
    "Bảng dưới mô tả hành vi đang áp dụng. Một lượt có giá cơ sở P = 120.000 đồng/khách; giá cuối còn nhân với số ghế tính tiền.",
)
case_rows = [
    ("A1", "Đi 24; không đặt về", "1 lượt đi ngày 24", "P. Bắt buộc chọn khung giờ."),
    ("A2", "Đi 25; không đặt về", "1 lượt đi ngày 25", "P. Bắt buộc chọn khung giờ."),
    ("A3", "Đi 24 và 25; không đặt về", "2 lượt đi: ngày 24 và ngày 25", "2P. Dùng cùng khung giờ đã chọn trong demo."),
    ("B1", "Đi 24; về 24", "Khứ hồi trong ngày 24", "2P x 90% = 1,8P. Tự chuyển Khứ hồi và hiện popup."),
    ("B2", "Đi 24; về 25", "1 lượt đi 24 + 1 lượt về 25", "2P. Tính hai lượt thường, không giảm 10%."),
    ("B3", "Đi 25; về 25", "Khứ hồi trong ngày 25", "2P x 90% = 1,8P. Tự chuyển Khứ hồi và hiện popup."),
    ("B4", "Đi 24 và 25; chỉ về 25", "Chuẩn hóa thành đi 24 + về 25; không có lượt đi 25", "2P. Hiện popup giải thích để tránh hiểu nhầm là 3 lượt."),
    ("B5", "Đi 24 và 25; về 24 và 25", "2 cặp khứ hồi: 24/24 và 25/25", "4P x 90% = 3,6P. Tự chuyển Khứ hồi và hiện popup."),
    ("B6", "Đi 24 và 25; chỉ về 24", "Đi 24 + về 24 + đi 25", "3P theo logic demo. Cần Product xác nhận trước khi phát hành thật."),
    ("C1", "Chỉ về ngày 24", "1 lượt từ Mỹ Đình về điểm người dùng chọn", "P. Không chọn giờ; xe đón sau khi concert kết thúc."),
    ("C2", "Chỉ về ngày 25", "1 lượt từ Mỹ Đình về điểm người dùng chọn", "P. Không chọn giờ; xe đón sau khi concert kết thúc."),
]
add_table(
    doc,
    ["Case", "Lựa chọn", "Các lượt sau chuẩn hóa", "Giá và phản hồi UI"],
    case_rows,
    [650, 2150, 3000, 3560],
    font_size=8.8,
)

add_heading(doc, "5.1. Ràng buộc chọn ngày", 2)
add_bullet(doc, "Nếu chỉ chọn một ngày đi, người dùng chỉ được chọn một ngày về, không được chọn đồng thời cả 24 và 25.")
add_bullet(doc, "Ngày về không được trước ngày đi sớm nhất. Ví dụ đi 25 thì ngày về 24 bị khóa.")
add_bullet(doc, "Chỉ khi chọn cả hai ngày đi mới cho phép chọn một hoặc cả hai ngày về.")
add_bullet(doc, "Nếu đổi ngày đi, hệ thống tự loại bỏ ngày về không còn hợp lệ và thu gọn về một ngày khi cần.")
add_bullet(doc, "Chọn trực tiếp Khứ hồi sẽ sao chép toàn bộ ngày đi sang ngày về.")
add_bullet(doc, "Chế độ Vé lượt về chỉ chọn đúng một ngày 24 hoặc 25.")

add_heading(doc, "6. Công thức giá hiện tại", 1)
add_text(doc, "Ký hiệu: P = 120.000 đồng; D = số ngày đi; R = số ngày về; S = số ghế tính tiền.")
add_table(
    doc,
    ["Điều kiện", "Giá một khách/ghế", "Công thức tổng"],
    [
        ("Vé lượt về", "P", "P x S"),
        ("Ngày đi và ngày về khớp hoàn toàn", "2P x 90% cho mỗi ngày", "(2P x 90% x D) x S"),
        ("Đi cả 24, 25 và chỉ về 25", "Hai lượt thường", "2P x S"),
        ("Các trường hợp thường khác", "P cho mỗi lượt", "P x (D + R) x S"),
        ("Bao xe 5/7 chỗ", "Giữ công thức hành trình", "S lần lượt bằng 5 hoặc 7, không phụ thuộc số người thực tế"),
    ],
    [2800, 2700, 3860],
    font_size=9.3,
)
add_callout(
    doc,
    "Lưu ý",
    "Mobile đang tính tạm để mô phỏng UI. Khi có backend, client gửi lựa chọn và hiển thị priceBreakdown do API quote trả về; không tự quyết định giá cuối để tránh lệch chính sách.",
    fill=LIGHT_GOLD,
    accent=GOLD,
)

add_heading(doc, "7. Logic khung giờ và popup giải thích", 1)
add_table(
    doc,
    ["Tình huống", "Hành vi UI"],
    [
        ("Vé lượt đi", "Hiện chọn khung giờ; bắt buộc chọn trước khi tạo vé. Nhắc có mặt trước giờ đi 15 phút."),
        ("Khứ hồi", "Vẫn chọn giờ lượt đi. Thời gian lượt về phụ thuộc lịch kết thúc concert và cần backend xác nhận."),
        ("Vé lượt về", "Ẩn toàn bộ chọn khung giờ. Ngay khi chọn loại vé, hiện popup: xe sẽ chờ đón sau khi concert kết thúc."),
        ("Đi/về trùng ngày", "Hiện popup giải thích hệ thống đã chuyển sang Khứ hồi và giảm 10% tổng hai lượt."),
        ("Đi cả hai ngày, chỉ về 25", "Hiện popup giải thích chỉ tính đi 24 và về 25, không phát sinh lượt đi 25."),
    ],
    [2650, 6710],
    font_size=9.5,
)

add_heading(doc, "8. Logic vị trí", 1)
add_table(
    doc,
    ["Loại hành trình", "Điểm đón", "Điểm đến/điểm trả"],
    [
        ("Lượt đi", "Người dùng tìm kiếm/chọn trên bản đồ bằng flow đặt chuyến hiện có", "Cố định: Sân vận động Quốc gia Mỹ Đình"),
        ("Khứ hồi", "Lượt đi dùng điểm người dùng chọn; lượt về xuất phát từ Mỹ Đình", "Lượt đi đến Mỹ Đình; lượt về quay về điểm đã chọn"),
        ("Lượt về", "Cố định: Sân vận động Quốc gia Mỹ Đình", "Người dùng tìm kiếm/chọn điểm trả bằng flow địa chỉ hiện có"),
    ],
    [1900, 3730, 3730],
    font_size=9.4,
)
add_bullet(doc, "Danh sách gợi ý địa chỉ và bản đồ tiếp tục dùng BookingModel/API chọn địa điểm của chức năng tạo chuyến cũ.")
add_bullet(doc, "Backend cần lưu snapshot địa chỉ và tọa độ tại thời điểm mua để không phụ thuộc kết quả tìm kiếm thay đổi về sau.")

add_heading(doc, "9. Vé, QR và trường hợp mua thêm vé", 1)
start_numbered_sequence(doc)
add_numbered(doc, "Tạo mới thay vì cập nhật", "Mỗi lần thanh toán thành công tạo một ticket/order mới. Vé cũ không biến mất và không bị thay QR.")
add_numbered(doc, "Điều hướng theo số lượng", "Không có vé: hiển thị trạng thái trống; một vé: mở thẳng chi tiết; từ hai vé: mở danh sách vé.")
add_numbered(doc, "Danh sách vé", "Mỗi thẻ hiển thị mã vé, ngày, hành trình, giờ/chính sách đón, điểm đón và trạng thái. Bấm thẻ để mở QR.")
add_numbered(doc, "Chi tiết vé", "QR luôn hiển thị ở đầu; không cần nút Xuất QR riêng. Bên dưới là ngày, giờ, hành trình, số khách, tổng tiền và lộ trình.")
add_numbered(doc, "Vé quá hạn", "Không xóa khỏi dữ liệu. Chuyển sang lịch sử hoặc trạng thái Đã hoàn thành/Hết hiệu lực; ưu tiên vé sắp dùng ở đầu danh sách.")
add_numbered(doc, "Nhiều ghế", "Bản demo dùng một QR cho một đơn và QR chứa số lượng khách. Backend cần chốt có phát hành một QR/đơn hay một QR/ghế trước khi làm app tài xế.")

add_heading(doc, "9.1. Dữ liệu tối thiểu của một vé", 2)
add_table(
    doc,
    ["Trường", "Mục đích"],
    [
        ("ticketId, ticketCode", "Định danh nội bộ và mã hiển thị cho khách hàng."),
        ("journeyType", "OUTBOUND_ONLY, RETURN_ONLY hoặc ROUND_TRIP; tránh suy luận từ ngày trống."),
        ("legs[]", "Danh sách lượt đã chuẩn hóa: chiều, serviceDate, địa chỉ, tọa độ, timeSlot/pickupPolicy."),
        ("vehicleType, isCharter, quantity", "Loại vé, bao xe và số ghế/khách cần kiểm."),
        ("priceBreakdown", "Giá cơ sở, số lượt, giảm giá, tổng tiền và mã chính sách."),
        ("status", "PENDING_PAYMENT, ACTIVE, PARTIALLY_USED, COMPLETED, CANCELLED, EXPIRED."),
        ("qrToken", "Token ký bởi server; không nhúng dữ liệu có thể giả mạo dưới dạng plain text trong production."),
        ("customer", "customerId; với khách mới kèm phone/email và kết quả cấp tài khoản."),
    ],
    [2400, 6960],
    font_size=9.4,
)

add_heading(doc, "10. Đề xuất flow API", 1)
start_numbered_sequence(doc)
add_numbered(doc, "Search location", "Mobile tái sử dụng API địa điểm hiện có để lấy địa chỉ và tọa độ.")
add_numbered(doc, "Quote", "Gửi hành trình người dùng chọn; server validate, chuẩn hóa legs và trả giá, discount, ruleCodes cùng thông điệp popup nếu cần.")
add_numbered(doc, "Create order", "Gửi quoteId và thông tin khách/loại xe. Dùng idempotencyKey để tránh tạo trùng khi người dùng bấm lại hoặc mạng chập chờn.")
add_numbered(doc, "Payment", "Server tạo giao dịch và trả QR thanh toán. Chỉ phát hành vé sau callback thanh toán thành công.")
add_numbered(doc, "Provision account", "Nếu là khách mới, server tạo tài khoản bằng số điện thoại và trả cơ chế đặt/nhận mật khẩu an toàn. Không để mobile tự sinh mật khẩu thật.")
add_numbered(doc, "Issue ticket", "Server trả ticket, qrToken và priceBreakdown chính thức; mobile mở màn hình chi tiết.")
add_numbered(doc, "List/detail", "GET tickets trả danh sách theo user/phone; mobile áp dụng quy tắc một vé mở thẳng, từ hai vé mở danh sách.")
add_numbered(doc, "Driver scan", "App tài xế gửi qrToken lên API verify; server trả hiệu lực, số khách còn lại và ghi nhận check-in theo idempotency.")

add_heading(doc, "10.1. Payload quote đề xuất", 2)
add_table(
    doc,
    ["Nhóm dữ liệu", "Ví dụ/ghi chú"],
    [
        ("journeyType", "OUTBOUND_ONLY | RETURN_ONLY | ROUND_TRIP"),
        ("outboundDates", "[2026-10-24, 2026-10-25] hoặc rỗng nếu chỉ về"),
        ("returnDates", "Một hoặc hai ngày theo ràng buộc nghiệp vụ"),
        ("outboundTimeSlot", "14:00/15:30/17:00/18:30; null với RETURN_ONLY"),
        ("pickup/dropoff", "placeId, address, latitude, longitude; server kiểm tra Mỹ Đình là đầu cố định phù hợp"),
        ("vehicle", "STANDARD_SEAT | CAR_5 | CAR_7, isCharter, quantity"),
        ("response", "normalizedLegs, priceBreakdown, ruleCodes, expiresAt, quoteId"),
    ],
    [2600, 6760],
    font_size=9.5,
)

add_heading(doc, "11. Validation và xử lý lỗi", 1)
add_bullet(doc, "Không cho tạo vé lượt đi/khứ hồi khi chưa chọn điểm đón hoặc chưa chọn khung giờ.")
add_bullet(doc, "Không cho tạo vé có lượt về khi danh sách ngày về rỗng.")
add_bullet(doc, "Không cho ngày về trước ngày đi sớm nhất.")
add_bullet(doc, "Khách mới phải có email và số điện thoại hợp lệ trước khi quote/create order.")
add_bullet(doc, "Server phải kiểm tra lại toàn bộ rule, tồn chỗ, giá và trạng thái quote; không tin dữ liệu tính giá từ client.")
add_bullet(doc, "Nếu quote hết hạn hoặc hết chỗ, trả lỗi có mã rõ ràng để mobile yêu cầu người dùng chọn lại.")
add_bullet(doc, "QR phải có thời hạn/chữ ký và verify online; không dùng mã QR do client tự sinh cho production.")

add_heading(doc, "12. Bộ test chấp nhận tối thiểu", 1)
test_rows = [
    ("T01", "Đi 24", "Tạo 1 lượt đi, giá P, bắt buộc chọn giờ"),
    ("T02", "Đi 24, về 24", "Tự chuyển khứ hồi, popup, giá 1,8P"),
    ("T03", "Đi 24, về 25", "Hai lượt thường, giá 2P"),
    ("T04", "Đi 25", "Ngày về 24 bị khóa"),
    ("T05", "Chọn một ngày đi", "Không thể chọn đồng thời hai ngày về"),
    ("T06", "Đi cả 24 và 25, chỉ về 25", "Chuẩn hóa đi 24 + về 25, giá 2P, có popup"),
    ("T07", "Đi cả 24 và 25, về cả 24 và 25", "Hai cặp khứ hồi, giá 3,6P, có popup"),
    ("T08", "Chỉ về 24/25", "Ẩn giờ, popup đón sau concert, giá P"),
    ("T09", "Bao xe 5/7 chỗ", "Khóa quantity theo 5/7 và tính giá theo toàn bộ sức chứa"),
    ("T10", "Tài khoản có một vé", "Mở thẳng chi tiết vé"),
    ("T11", "Tài khoản có hai vé", "Mở danh sách; chọn được từng QR"),
    ("T12", "Khách mới", "Validate phone/email; vé hiển thị thông tin đăng nhập được cấp"),
    ("T13", "Token hết hạn/đăng xuất", "Login được điền sẵn thông tin phiên cuối; X xóa cả phone/password"),
    ("T14", "Bấm tạo đơn lặp do timeout", "Backend idempotency không tạo hai vé"),
]
add_table(
    doc,
    ["ID", "Dữ liệu kiểm thử", "Kết quả mong đợi"],
    test_rows,
    [700, 2800, 5860],
    font_size=9.1,
)

add_heading(doc, "13. Mapping vào mã nguồn hiện tại", 1)
add_table(
    doc,
    ["Tệp", "Trách nhiệm"],
    [
        ("lib/screens/concert/concert_booking_screen.dart", "Form đặt vé, state ngày/hành trình, giá demo, tạo dữ liệu vé, danh sách vé và chi tiết QR."),
        ("lib/screens/popup/concert_round_trip_popup.dart", "Popup khứ hồi, case đi 24/về 25 và thông báo thời gian đón lượt về."),
        ("lib/screens/beluca_home_view.dart", "Entry point Mua vé/Xem vé tại Home và banner concert."),
        ("lib/screens/login_screen.dart", "Entry point khách chưa đăng nhập, tự điền và xóa thông tin đăng nhập đã lưu."),
        ("lib/services/login_credential_storage.dart", "Lưu/xóa phone và password phiên cuối bằng secure storage."),
        ("lib/models/booking_model.dart", "Tái sử dụng tìm kiếm, chọn và hiển thị địa chỉ từ flow đặt chuyến."),
    ],
    [3400, 5960],
    font_size=9.2,
)

page_break = doc.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)
add_heading(doc, "14. Điểm cần chốt trước khi tích hợp production", 1)
add_table(
    doc,
    ["Mức độ", "Nội dung cần quyết định"],
    [
        ("Cao", "Case B6: đi cả 24/25 nhưng chỉ về 24 có được phép không; nếu có thì xác nhận đây là 3 lượt."),
        ("Cao", "Một QR đại diện cho cả đơn nhiều khách hay cấp QR riêng cho từng ghế/hành khách."),
        ("Cao", "Thời điểm cụ thể và điểm tập kết của lượt về sau khi concert kết thúc; cách xử lý concert kết thúc muộn."),
        ("Cao", "Chính sách hủy/đổi ngày/đổi số lượng, hoàn tiền và giới hạn bán theo sức chứa."),
        ("Trung bình", "Khứ hồi khác ngày có được giảm giá hay chỉ giảm khi đi/về cùng ngày như demo."),
        ("Trung bình", "Một khung giờ áp dụng cho cả hai ngày hay người dùng phải chọn giờ riêng cho từng ngày."),
        ("Trung bình", "Cách cấp mật khẩu khách mới: mật khẩu tạm, OTP đặt mật khẩu hay magic link; không trả mật khẩu dài hạn dạng plain text."),
        ("Trung bình", "Quy tắc sắp xếp/lưu trữ vé đã dùng và thời điểm chuyển ACTIVE sang COMPLETED/EXPIRED."),
    ],
    [1300, 8060],
    font_size=9.3,
)

add_callout(
    doc,
    "Kết luận bàn giao",
    "UI demo đã bao phủ các flow chính và các case ngày quan trọng. Bước tiếp theo nên là chốt các điểm mở ở Mục 14, sau đó thiết kế API quote trả normalizedLegs + priceBreakdown để dùng chung cho Mobile, Backend và QA.",
    fill=LIGHT_GREEN,
    accent=GREEN,
)

# Document metadata.
doc.core_properties.title = "Đặc tả flow đặt vé xe concert"
doc.core_properties.subject = "Nghiệp vụ, logic ngày đi/ngày về và đề xuất API"
doc.core_properties.author = "Belucar Product & Engineering"
doc.core_properties.keywords = "Belucar, concert, ticket, booking, API, QR"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
